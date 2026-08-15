"""
Robi "zloty zestaw testowy" z pliku w formacie MasterTemplate 5.17.

Wyciaga jednoczesnie WEJSCIE (co trener wpisal) i WYNIK (co arkusz policzyl),
zeby dalo sie sprawdzic, czy silnik liczy to samo.

Uzycie:
    python3 silnik/narzedzia/zrzut-arkusza.py plan.xlsx silnik/testy/zlote/nazwa.json

WAZNE: plik musi byc przeliczony. Arkusz zapisuje ostatnio wyliczone wartosci —
jesli byl edytowany bez przeliczenia, czesc komorek bedzie pusta i te pola
zostana pominiete w porownaniu (skrypt to raportuje).
"""
import json, sys, pathlib

try:
    import openpyxl
except ImportError:
    sys.exit("Brak openpyxl. Zainstaluj: apt-get install -y python3-openpyxl")

TYGODNIE = ["T1", "T2", "T3", "T4", "T5", "T6"]
LICZBA_DNI = 5
SLOTOW_W_DNIU = 12

# Kolumny w zakladkach T1-T6 (1 = A)
KOL = {"lp": 1, "przelacznik": 2, "cwiczenie": 3, "serie": 4, "powt": 5, "rpe": 6,
       "ciezar": 7, "feedback": 8, "szkielet": 10, "position": 11, "ex_id": 12,
       "part": 13, "kategoria": 14, "coeff": 15, "skok": 16, "progresja": 17,
       "one_rm": 18, "procent": 19, "stres_t": 20, "stres_c": 21, "stres_p": 22,
       "one_rm_reczny": 27, "mnoznik": 29, "korekta_powt": 30}

# Kolumny w zakladce START
KOL_START = {"lp": 1, "cwiczenie": 2, "ciezar": 3, "powt": 4, "one_rm": 5, "position": 11}


def wiersz_topsetu(dzien: int) -> int:
    return 6 + (dzien - 1) * 16


def wiersze_slotow(dzien: int) -> range:
    start = wiersz_topsetu(dzien) + 1
    return range(start, start + SLOTOW_W_DNIU)


def wiersz_podsumowania(dzien: int) -> int:
    return wiersz_topsetu(dzien) + 13


def wiersze_startu(dzien: int) -> range:
    start = 6 + (dzien - 1) * 13
    return range(start, start + SLOTOW_W_DNIU)


def liczba(v):
    return v if isinstance(v, (int, float)) else None


def tekst(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def zrzut(sciezka_xlsx: str) -> dict:
    wb = openpyxl.load_workbook(sciezka_xlsx, data_only=True)
    braki = {"komorek_bez_wartosci": 0}

    analiza = wb["Analiza"]
    tryb = tekst(analiza["B4"].value) or "trzymaj z bloku"
    czesc_raw = tekst(analiza["B5"].value) or "część 1 (objętość)"
    czesc = "objętość" if "objęto" in czesc_raw else "intensywność"

    # --- serie maksymalne ---
    # Kolumna B w START to formula lustrzana z T1 — w nieprzeliczonym pliku bywa pusta.
    # position_id (kolumna K) jest wartoscia stala, wiec wiazemy po nim.
    ws = wb["START"]
    t1 = wb["T1"]
    serie_max = []
    for dzien in range(1, LICZBA_DNI + 1):
        for r in wiersze_startu(dzien):
            ciezar = liczba(ws.cell(row=r, column=KOL_START["ciezar"]).value)
            powt = liczba(ws.cell(row=r, column=KOL_START["powt"]).value)
            if ciezar is None or powt is None:
                continue
            serie_max.append({
                "position_id": tekst(ws.cell(row=r, column=KOL_START["position"]).value),
                "nazwa": tekst(ws.cell(row=r, column=KOL_START["cwiczenie"]).value),
                "ciezar": ciezar,
                "powtorzenia": powt,
                "oczekiwany_1rm": liczba(ws.cell(row=r, column=KOL_START["one_rm"]).value),
            })

    # --- sloty i wyniki per tydzien ---
    sloty = {}
    for dzien in range(1, LICZBA_DNI + 1):
        for r in wiersze_slotow(dzien):
            nazwa = tekst(t1.cell(row=r, column=KOL["cwiczenie"]).value)
            if not nazwa:
                continue
            pid = tekst(t1.cell(row=r, column=KOL["position"]).value)
            sloty[pid] = {
                "position_id": pid,
                "dzien": dzien,
                "wiersz": r,
                "lp": tekst(t1.cell(row=r, column=KOL["lp"]).value),
                "nazwa": nazwa,
                "ex_id": tekst(t1.cell(row=r, column=KOL["ex_id"]).value),
                "kategoria_szkieletu": tekst(t1.cell(row=r, column=KOL["szkielet"]).value),
                "tygodnie": {},
            }

    for skrot in TYGODNIE:
        ws = wb[skrot]
        for pid, slot in sloty.items():
            r = slot["wiersz"]
            pole = {
                "serie": liczba(ws.cell(row=r, column=KOL["serie"]).value),
                "rpe": liczba(ws.cell(row=r, column=KOL["rpe"]).value),
                "feedback": tekst(ws.cell(row=r, column=KOL["feedback"]).value),
                "cwiczenie": tekst(ws.cell(row=r, column=KOL["cwiczenie"]).value),
                "one_rm_reczny": liczba(ws.cell(row=r, column=KOL["one_rm_reczny"]).value),
                # oczekiwane wyniki
                "ocz_powtorzenia": liczba(ws.cell(row=r, column=KOL["powt"]).value),
                "ocz_ciezar": ws.cell(row=r, column=KOL["ciezar"]).value,
                "ocz_procent": liczba(ws.cell(row=r, column=KOL["procent"]).value),
                "ocz_one_rm": liczba(ws.cell(row=r, column=KOL["one_rm"]).value),
                "ocz_mnoznik": liczba(ws.cell(row=r, column=KOL["mnoznik"]).value),
                "ocz_stres_t": liczba(ws.cell(row=r, column=KOL["stres_t"]).value),
                "ocz_stres_c": liczba(ws.cell(row=r, column=KOL["stres_c"]).value),
                "ocz_stres_p": liczba(ws.cell(row=r, column=KOL["stres_p"]).value),
            }
            for k, v in pole.items():
                if k.startswith("ocz_") and v is None:
                    braki["komorek_bez_wartosci"] += 1
            slot["tygodnie"][skrot] = pole

    # --- TOP SETY ---
    top_sety = []
    for dzien in range(1, LICZBA_DNI + 1):
        r = wiersz_topsetu(dzien)
        wlaczony = tekst(t1.cell(row=r, column=KOL["przelacznik"]).value) == "TOP SET"
        pierwszy = tekst(t1.cell(row=r + 1, column=KOL["position"]).value)
        top_sety.append({
            "dzien": dzien,
            "wlaczony": wlaczony,
            "rpe": liczba(t1.cell(row=r, column=KOL["rpe"]).value),
            "slot_position_id": pierwszy,
        })

    # --- podsumowania dni i bilans tygodnia ---
    podsumowania = {}
    for skrot in TYGODNIE:
        ws = wb[skrot]
        dni = []
        for dzien in range(1, LICZBA_DNI + 1):
            r = wiersz_podsumowania(dzien)
            dni.append({
                "dzien": dzien,
                "ocz_serie": liczba(ws.cell(row=r, column=4).value),
                "ocz_powtorzenia": tekst(ws.cell(row=r, column=5).value),
            })
        wzorce = {}
        for i, part in enumerate(["s", "d", "b", "r", "c"]):
            r = 87 + i
            wzorce[part] = {
                "ocz_calkowity": liczba(ws.cell(row=r, column=4).value),
                "ocz_centralny": liczba(ws.cell(row=r, column=6).value),
                "ocz_obwodowy": liczba(ws.cell(row=r, column=7).value),
                "ocz_serie": liczba(ws.cell(row=r, column=8).value),
                "ocz_powtorzenia": liczba(ws.cell(row=r, column=9).value),
            }
        podsumowania[skrot] = {
            "dni": dni,
            "wzorce": wzorce,
            "ocz_razem": liczba(ws.cell(row=94, column=4).value),
            "ocz_serie_razem": liczba(ws.cell(row=94, column=8).value),
            "ocz_powtorzenia_razem": liczba(ws.cell(row=94, column=9).value),
        }

    # Powiazanie serii maksymalnych z cwiczeniem po position_id slotu w T1.
    for s in serie_max:
        slot = sloty.get(s["position_id"])
        s["ex_id"] = slot["ex_id"] if slot else None
        if slot and not s["nazwa"]:
            s["nazwa"] = slot["nazwa"]

    # Slot, w ktorym arkusz nie ma wyliczonego 1RM mimo istniejacej serii maksymalnej,
    # znaczy tylko tyle, ze plik nie byl przeliczony — porownanie ciezaru trzeba pominac.
    z_seria = {s["ex_id"] for s in serie_max if s.get("ex_id")}
    for slot in sloty.values():
        t1_pole = slot["tygodnie"]["T1"]
        slot["arkusz_nieprzeliczony"] = bool(
            slot["ex_id"] in z_seria and t1_pole["ocz_one_rm"] is None
        )
        if slot["arkusz_nieprzeliczony"]:
            braki.setdefault("sloty_z_nieprzeliczonym_1rm", []).append(slot["position_id"])

    return {
        "zrodlo": pathlib.Path(sciezka_xlsx).name,
        "ustawienia": {"tryb_akcesoriow": tryb, "czesc_planu": czesc,
                       "dni_treningowe": liczba(analiza["B76"].value)},
        "serie_maksymalne": serie_max,
        "sloty": list(sloty.values()),
        "top_sety": top_sety,
        "podsumowania": podsumowania,
        "diagnostyka": braki,
    }


if __name__ == "__main__":
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    dane = zrzut(sys.argv[1])
    pathlib.Path(sys.argv[2]).write_text(
        json.dumps(dane, ensure_ascii=False, indent=1), encoding="utf-8")
    print(f"{dane['zrodlo']}: {len(dane['sloty'])} slotow, "
          f"{len(dane['serie_maksymalne'])} serii maksymalnych, "
          f"{dane['diagnostyka']['komorek_bez_wartosci']} komorek bez wartosci "
          f"(pominiete w porownaniu)")
    print(f"-> {sys.argv[2]}")
